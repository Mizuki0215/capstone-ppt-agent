from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- palette ----------
INK     = RGBColor(0x10, 0x23, 0x3A)
INK2    = RGBColor(0x1E, 0x36, 0x52)
GOLD    = RGBColor(0xC0, 0x8A, 0x2E)
TEAL    = RGBColor(0x2E, 0x6F, 0x95)
RED     = RGBColor(0xA8, 0x35, 0x42)
GREY    = RGBColor(0x5A, 0x6B, 0x7C)
LGREY   = RGBColor(0x8A, 0x99, 0xA8)
PANEL   = RGBColor(0xF4, 0xF6, 0xF8)
PANEL2  = RGBColor(0xEC, 0xF1, 0xF5)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
REDBG   = RGBColor(0xFA, 0xEE, 0xEF)
GOLDBG  = RGBColor(0xF9, 0xF3, 0xE4)

FONT = "PingFang HK"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def new_slide():
    return prs.slides.add_slide(BLANK)


def nostyle(sh):
    el = sh._element
    st = el.find(qn('p:style'))
    if st is not None:
        el.remove(st)
    return sh


def rrect(s, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.035, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    nostyle(sh)
    return sh


def tbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0
    tf.margin_top = 0; tf.margin_bottom = 0
    return tf


def add(tf, text, size=13, bold=False, color=INK, align=PP_ALIGN.LEFT,
        space_before=0, space_after=0, line=1.18, italic=False, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run(); r.text = text
    f = r.font
    f.name = FONT; f.size = Pt(size); f.bold = bold; f.italic = italic; f.color.rgb = color
    return p


def header(s, kicker, title, sub=None):
    tf = tbox(s, 0.85, 0.42, 11.6, 0.3)
    add(tf, kicker.upper(), size=9.5, bold=True, color=GOLD, first=True)
    tf = tbox(s, 0.85, 0.72, 11.6, 0.62)
    add(tf, title, size=27, bold=True, color=INK, first=True)
    if sub:
        tf = tbox(s, 0.85, 1.42, 11.6, 0.35)
        add(tf, sub, size=12, color=GREY, first=True)
    ln = s.shapes.add_connector(1, Inches(0.85), Inches(1.32 if not sub else 1.86),
                                Inches(12.48), Inches(1.32 if not sub else 1.86))
    ln.line.color.rgb = RGBColor(0xDD, 0xE2, 0xE8); ln.line.width = Pt(0.9)


def footer(s, page, dark=False):
    c = LGREY if not dark else RGBColor(0x8F, 0xA3, 0xB5)
    tf = tbox(s, 0.85, 7.02, 8.0, 0.25)
    add(tf, "DSAI4901  ·  Capstone Project Proposal  ·  CHAN Yee Tung", size=8, color=c, first=True)
    tf = tbox(s, 12.0, 7.02, 0.5, 0.25)
    add(tf, f"{page:02d}", size=8.5, bold=True, color=c, align=PP_ALIGN.RIGHT, first=True)


def bullets(tf, items, size=12, color=INK2, gap=5, line=1.25, mark="•", mark_color=GOLD):
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else gap)
        p.line_spacing = line
        r1 = p.add_run(); r1.text = mark + "  "
        r1.font.name = FONT; r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = mark_color
        r2 = p.add_run(); r2.text = it
        r2.font.name = FONT; r2.font.size = Pt(size); r2.font.color.rgb = color


# =====================================================================
# 01 TITLE
# =====================================================================
s = new_slide()
rrect(s, 0, 0, 0.28, H, fill=INK, shape=MSO_SHAPE.RECTANGLE)
rrect(s, 0.28, 0, 0.06, H, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)

tf = tbox(s, 1.0, 1.30, 8.0, 0.3)
add(tf, "DSAI4901  ·  CAPSTONE PROJECT PROPOSAL", size=10, bold=True, color=GOLD, first=True)

tf = tbox(s, 1.0, 1.85, 8.3, 1.9)
add(tf, "論文整理型 PPT Agent", size=44, bold=True, color=INK, line=1.05, first=True)
add(tf, "從「生成得到」到「驗證得到」", size=30, bold=False, color=TEAL, space_before=10, line=1.1)

rrect(s, 1.0, 4.05, 1.5, 0.045, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)

tf = tbox(s, 1.0, 4.30, 7.6, 0.8)
add(tf, "以內容可追溯性（traceability）與版面可用率（deck-readiness）\n為核心的簡報生成系統",
    size=13.5, color=GREY, line=1.45, first=True)

tf = tbox(s, 1.0, 5.55, 5.5, 1.2)
add(tf, "CHAN Yee Tung", size=13, bold=True, color=INK, first=True)
add(tf, "Supervisor：Mr. Chen Xianwei", size=10.5, color=GREY, space_before=3)
add(tf, "2026 年 10 月", size=10.5, color=GREY, space_before=1)

# right stats
stats = [("35", "篇核心文獻"), ("8", "個功能缺口"), ("1", "條研究路線")]
for i, (num, lab) in enumerate(stats):
    y = 1.95 + i * 1.62
    tf = tbox(s, 9.75, y, 2.6, 0.8)
    add(tf, num, size=46, bold=True, color=GOLD, line=0.95, first=True)
    tf = tbox(s, 9.75, y + 0.72, 2.6, 0.3)
    add(tf, lab, size=10.5, color=GREY, first=True)
    if i < 2:
        ln = s.shapes.add_connector(1, Inches(9.75), Inches(y + 1.42), Inches(12.3), Inches(y + 1.42))
        ln.line.color.rgb = RGBColor(0xE4, 0xE8, 0xEC); ln.line.width = Pt(0.75)

footer(s, 1)

# =====================================================================
# 02 LITERATURE SCALE
# =====================================================================
s = new_slide()
header(s, "01 · 研究基礎", "文獻研究：由 367 篇收窄到 35 篇")

rows = [("367", "arXiv 候選", 8.6, INK),
        ("69",  "相關度篩選後", 5.6, INK2),
        ("35",  "精讀（30 Agent + 5 教育）", 4.2, TEAL),
        ("18",  "作者公開原始碼", 3.6, GOLD)]
y0 = 2.10
for i, (num, lab, wbar, col) in enumerate(rows):
    y = y0 + i * 0.98
    rrect(s, 0.85, y, wbar, 0.74, fill=col if i == 3 else PANEL, line=None)
    tf = tbox(s, 0.85 + 0.22, y + 0.13, 0.95, 0.5)
    add(tf, num, size=21, bold=True, color=WHITE if i == 3 else INK, first=True)
    tf = tbox(s, 0.85 + 1.25, y + 0.20, wbar - 1.3, 0.36)
    add(tf, lab, size=11, color=WHITE if i == 3 else GREY, first=True)

# right method box
rrect(s, 9.35, 2.10, 3.1, 3.30, fill=PANEL2)
tf = tbox(s, 9.60, 2.28, 2.62, 0.3)
add(tf, "我實際做了什麼", size=10, bold=True, color=GOLD, first=True)
tf = tbox(s, 9.60, 2.62, 2.62, 3.3)
bullets(tf, ["12 組 arXiv 查詢（slide / PPT / poster / deck）",
             "以功能相關度篩選，排除無關領域",
             "35 篇全文關鍵詞掃描（WCAG、notes、成本）",
             "逐篇由 PDF 抽取 GitHub 連結並 HTTP 驗證"],
        size=9.5, gap=7, line=1.28)

tf = tbox(s, 0.85, 6.25, 8.2, 0.4)
add(tf, "18/35 有作者公開原始碼（已逐一 HTTP 驗證）；其餘為論文描述或匿名投稿",
    size=9, color=LGREY, first=True)
footer(s, 2)

# =====================================================================
# 03 WHAT THEY DO WELL
# =====================================================================
s = new_slide()
header(s, "01 · 研究基礎", "佢哋做得好嘅四件事", "先確認：現有系統已解決了什麼——這決定了我不應該重做什麼")

cards = [
    ("01", "生成品質與版面", "已做到「看起來像真」：版面穩定、風格一致，盲測可贏商業工具。",
     "Zheng et al. (2025); Vats et al. (2026)"),
    ("02", "原生可編輯輸出", "開始在 .pptx 物件層操作，保留樣式與可編輯性。",
     "Jung et al. (2025); Zheng et al. (2026b)"),
    ("03", "評測方法", "從主觀打分走向任務化：用答題測內容傳達、用 gate 測可用性。",
     "Pang et al. (2025); Hu et al. (2026)"),
    ("04", "多模態與敘事", "圖、表、公式能進入版面；開始顯式建模故事線。",
     "Liang et al. (2025); Ozden et al. (2026a)"),
]
cw, gap = 2.86, 0.22
for i, (num, title, body, cite) in enumerate(cards):
    x = 0.85 + i * (cw + gap)
    rrect(s, x, 2.15, cw, 3.20, fill=PANEL)
    rrect(s, x, 2.15, cw, 0.05, fill=GOLD, shape=MSO_SHAPE.RECTANGLE)
    tf = tbox(s, x + 0.24, 2.38, cw - 0.4, 0.35)
    add(tf, num, size=11, bold=True, color=GOLD, first=True)
    tf = tbox(s, x + 0.24, 2.72, cw - 0.42, 0.7)
    add(tf, title, size=14.5, bold=True, color=INK, line=1.15, first=True)
    tf = tbox(s, x + 0.24, 3.35, cw - 0.42, 1.15)
    add(tf, body, size=10.5, color=INK2, line=1.38, first=True)
    tf = tbox(s, x + 0.24, 4.55, cw - 0.42, 0.5)
    add(tf, cite, size=8, color=TEAL, line=1.2, first=True)

tf = tbox(s, 0.85, 5.85, 11.6, 0.5)
add(tf, "結論：生成能力已經商品化，我不應該再做一個更好的生成器。",
    size=12, bold=True, color=RED, first=True)
footer(s, 3)

# =====================================================================
# 04 FUNCTION MAP
# =====================================================================
s = new_slide()
header(s, "01 · 研究基礎", "把 PPT Agent 拆成 8 個功能", "缺口要講得具體，就先要有一張功能地圖——之後四頁逐個對照")

funcs = [("①", "解析"), ("②", "內容選取"), ("③", "敘事"), ("④", "版面"), ("⑤", "輸出"), ("⑥", "檢查")]
bw, bh, bgap = 1.62, 1.30, 0.30
x0 = 0.85
for i, (num, name) in enumerate(funcs):
    x = x0 + i * (bw + bgap)
    rrect(s, x, 2.30, bw, bh, fill=PANEL2, line=RGBColor(0xD8, 0xE0, 0xE8))
    tf = tbox(s, x, 2.46, bw, 0.32)
    add(tf, num, size=13, bold=True, color=GOLD, align=PP_ALIGN.CENTER, first=True)
    tf = tbox(s, x, 2.86, bw, 0.4)
    add(tf, name, size=13.5, bold=True, color=INK, align=PP_ALIGN.CENTER, first=True)
    if i < 5:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.05), Inches(2.83),
                                Inches(0.20), Inches(0.24))
        ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor(0xC5, 0xCF, 0xD9)
        ar.line.fill.background(); ar.shadow.inherit = False; nostyle(ar)

# cross-cutting
for i, (num, name, desc) in enumerate([
        ("⑦", "講稿", "跨階段：把內容轉成可講的節奏"),
        ("⑧", "引用與可追溯", "跨階段：每一句都要有出處")]):
    y = 4.15 + i * 1.05
    rrect(s, 0.85, y, 11.6, 0.86, fill=GOLDBG, line=None)
    tf = tbox(s, 1.10, y + 0.10, 0.5, 0.6)
    add(tf, num, size=15, bold=True, color=GOLD, first=True)
    tf = tbox(s, 1.62, y + 0.11, 2.4, 0.6)
    add(tf, name, size=13.5, bold=True, color=INK, first=True)
    tf = tbox(s, 4.10, y + 0.20, 8.0, 0.5)
    add(tf, desc, size=11, color=GREY, first=True)

tf = tbox(s, 0.85, 6.55, 11.6, 0.35)
add(tf, "接下來四頁：每頁講 2 個缺口 —— 症狀 · 文獻證據 · 後果",
    size=10.5, color=LGREY, first=True)
footer(s, 4)


# ---------- weakness slide helper ----------
def weakness_slide(page, kicker, title, left, right):
    s = new_slide()
    header(s, kicker, title)
    for idx, d in enumerate([left, right]):
        x = 0.85 + idx * 6.02
        wcol = 5.72
        rrect(s, x, 1.72, wcol, 4.95, fill=WHITE, line=RGBColor(0xDF, 0xE5, 0xEB), lw=0.9)
        rrect(s, x, 1.72, wcol, 0.055, fill=RED, shape=MSO_SHAPE.RECTANGLE)

        tf = tbox(s, x + 0.28, 1.92, wcol - 0.5, 0.36)
        add(tf, d["label"], size=10, bold=True, color=RED, first=True)
        tf = tbox(s, x + 0.28, 2.22, wcol - 0.5, 0.42)
        add(tf, d["name"], size=17.5, bold=True, color=INK, first=True)

        tf = tbox(s, x + 0.28, 2.76, wcol - 0.5, 0.22)
        add(tf, "症狀", size=8.5, bold=True, color=LGREY, first=True)
        tf = tbox(s, x + 0.28, 2.97, wcol - 0.5, 1.35)
        bullets(tf, d["symptoms"], size=10, gap=3, line=1.22, mark="▪", mark_color=RED)

        tf = tbox(s, x + 0.28, 4.14, wcol - 0.5, 0.22)
        add(tf, "文獻證據 · APA", size=8.5, bold=True, color=LGREY, first=True)
        tf = tbox(s, x + 0.28, 4.35, wcol - 0.5, 0.85)
        add(tf, d["evidence"], size=9, color=TEAL, line=1.3, first=True)

        rrect(s, x + 0.20, 5.16, wcol - 0.40, 1.12, fill=REDBG, line=None)
        tf = tbox(s, x + 0.36, 5.25, wcol - 0.70, 0.2)
        add(tf, "後果", size=8.5, bold=True, color=RED, first=True)
        tf = tbox(s, x + 0.36, 5.46, wcol - 0.70, 0.74)
        add(tf, d["consequence"], size=9.8, color=INK, line=1.28, first=True)
    footer(s, page)
    return s


# =====================================================================
# 05 GAP 1-2
# =====================================================================
weakness_slide(5, "02 · 功能缺口", "缺口 1 ｜ 解析與內容選取",
    {"label": "① 解析", "name": "圖文分離，輸入即失真",
     "symptoms": ["figure 與 caption 被拆散，圖文對不上",
                  "雙欄論文閱讀順序錯亂，段落拼接錯位",
                  "表格被壓成文字牆，數字失去行列結構"],
     "evidence": "現有系統把解析視為前處理，未處理圖文對齊（Pang et al., 2025; Fu et al., 2021）。",
     "consequence": "錯的輸入會被下游全部繼承——後面每一步都在錯的基礎上做決策，而且錯誤在成品上不可見。"},
    {"label": "② 內容選取", "name": "平均分配，沒有輕重",
     "symptoms": ["每節平均分配一頁，篇幅與重要性無關",
                  "最關鍵的實驗結果被省略，反而保留背景",
                  "method 細節塞進 introduction 頁"],
     "evidence": "多數方法集中於文字選取，忽略篇幅與重要性的對應（Liang et al., 2025; Chen et al., 2026）。",
     "consequence": "deck 完整但沒有重點；觀眾離場只記得標題，講者也說不出「最重要的一句是什麼」。"})

# =====================================================================
# 06 GAP 3-4
# =====================================================================
weakness_slide(6, "02 · 功能缺口", "缺口 2 ｜ 敘事與版面",
    {"label": "③ 敘事", "name": "目錄式摘要，不是演講",
     "symptoms": ["逐段摘要＝論文目錄，不是一條論證線",
                  "10 頁 background、2 頁 result，輕重倒置",
                  "缺少 motivation → problem → insight → evidence 的弧線"],
     "evidence": "現有方法直接把原文摘要成 slide（Ozden et al., 2026a）；顯式敘事規劃仍屬少數（Vats et al., 2026）。",
     "consequence": "講者變成朗讀目錄，聽眾無法建立因果理解——記得住每頁，但說不出整個研究在解決什麼。"},
    {"label": "④ 版面", "name": "出界、重疊、被裁切",
     "symptoms": ["文字超出版面、元素互相重疊",
                  "字級過小、對比不足，投影後看不清",
                  "圖被裁切、留白失衡，每頁都是 bullet 牆"],
     "evidence": "局部錯誤包括 overflow、overlap、clipping、off-canvas（Tian et al., 2026）；模型讀得懂內容卻排不好版面（Huang et al., 2025）。",
     "consequence": "這是最直接可見的失敗：投影當場文字被切、圖被遮，觀眾即時失去信任。"})

# =====================================================================
# 07 GAP 5-6
# =====================================================================
weakness_slide(7, "02 · 功能缺口", "缺口 3 ｜ 輸出與檢查",
    {"label": "⑤ 輸出", "name": "看似可用，其實不能改",
     "symptoms": ["輸出圖片或 HTML，無法在 PowerPoint 內編輯",
                  "出到 .pptx 但全是 textbox，改一個字就爆版",
                  "圖表被點陣化，原始資料消失"],
     "evidence": "現有 agent 常把可編輯內容點陣化或破壞版面（Zheng et al., 2026b）；物件模型路線才能保留樣式（Jung et al., 2025）。",
     "consequence": "用戶拿到成品但不能改，等於不能用；修改成本高於自己重做，工具因此被棄用。"},
    {"label": "⑥ 檢查", "name": "一次性回饋，無法歸因",
     "symptoms": ["整份重寫一次，只看一次回饋",
                  "錯誤無法定位到具體元素（哪個 shape 出界？）",
                  "修完可能更差，沒有收斂保證"],
     "evidence": "單體式 one-version-one-feedback 令局部錯誤難以歸因（Tian et al., 2026）；反思須以渲染狀態為依據（Zheng et al., 2026a）。",
     "consequence": "迭代變成賭博：越修越差，而且無法解釋為什麼——使用者最終選擇手動修改。"})

# =====================================================================
# 08 GAP 7-8
# =====================================================================
weakness_slide(8, "02 · 功能缺口", "缺口 4 ｜ 講稿與引用",
    {"label": "⑦ 講稿", "name": "notes 只是 bullet 的複述",
     "symptoms": ["speaker notes 與頁面文字高度重複",
                  "沒有時間預算，不知 3 分鐘該講幾頁",
                  "沒有轉場句，頁與頁之間斷開"],
     "evidence": "現有系統低估交付過程（Yang et al., 2026）；35 篇中僅 5 篇提及 speaker notes，且均非研究重點。",
     "consequence": "講者照讀 slide，觀眾同時讀字與聽字，注意力分裂；準備時間沒有減少，工具失去意義。"},
    {"label": "⑧ 引用與可追溯", "name": "寫得出字，說不出出處",
     "symptoms": ["bullet 沒有出處，無法回溯到論文哪一段",
                  "把 related work 寫成「本文提出」",
                  "圖表編號錯配，數字張冠李戴"],
     "evidence": "檢索用於提升可靠性但未提供逐句出處（Xie et al., 2025）；全局意圖有保留，句子級錨點仍缺（Ozden et al., 2026a）。",
     "consequence": "講錯數據也無人察覺；在學術場景，一次錯誤引用就足以摧毀整場簡報的可信度。"})

# =====================================================================
# 09 CONSEQUENCE CHAIN
# =====================================================================
s = new_slide()
header(s, "02 · 功能缺口", "單點失效，會被下游放大")

chain = ["輸入錯", "內容錯", "敘事散", "版面爆", "修不好", "無從查證"]
cw2, ch2, gap2 = 1.72, 0.72, 0.26
x0 = 0.85
for i, t in enumerate(chain):
    x = x0 + i * (cw2 + gap2)
    rrect(s, x, 1.85, cw2, ch2, fill=INK if i < 5 else RED, line=None)
    tf = tbox(s, x, 1.85 + 0.17, cw2, 0.4)
    add(tf, t, size=11.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
    if i < 5:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + cw2 + 0.04), Inches(2.12),
                                Inches(0.18), Inches(0.18))
        ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor(0xC5, 0xCF, 0xD9)
        ar.line.fill.background(); ar.shadow.inherit = False; nostyle(ar)

tf = tbox(s, 0.85, 2.72, 11.6, 0.3)
add(tf, "每一個上游錯誤，都會在下游被放大，而且無法回溯", size=10, color=LGREY, first=True)

impacts = [("觀眾", "看不懂、記不住，當場看到字被切", "注意力流失"),
           ("講者", "不敢用、要重做，答不出「這數字哪來的」", "信任崩塌"),
           ("研究者", "無法比較系統好壞，因為沒有共同的可信度量", "領域停滯")]
for i, (who, what, tag) in enumerate(impacts):
    x = 0.85 + i * 3.92
    rrect(s, x, 3.35, 3.66, 1.60, fill=PANEL)
    tf = tbox(s, x + 0.28, 3.52, 3.1, 0.3)
    add(tf, who, size=12, bold=True, color=GOLD, first=True)
    tf = tbox(s, x + 0.28, 3.86, 3.1, 1.0)
    add(tf, what, size=10.2, color=INK2, line=1.4, first=True)
    tf = tbox(s, x + 0.28, 4.60, 3.1, 0.3)
    add(tf, "→ " + tag, size=9.5, bold=True, color=RED, first=True)

rrect(s, 0.85, 5.62, 11.6, 0.95, fill=GOLDBG, line=None)
tf = tbox(s, 1.15, 5.78, 11.0, 0.65)
add(tf, "35 篇的核心問題：全部在優化「生成」，沒有一篇在量度「可信」。",
    size=14, bold=True, color=INK, first=True)
footer(s, 9)

# =====================================================================
# 10 MY DIRECTION
# =====================================================================
s = new_slide()
header(s, "03 · 研究定位", "我的方向：論文整理型 PPT Agent", "明確不是 lecture note —— 兩者的成功標準根本不同")

rows = [("輸入", "單篇論文 PDF", "教材 / 課堂內容"),
        ("目標", "把研究壓縮成可講的論證鏈", "把知識拆成學習單元"),
        ("成功標準", "內容可追溯、論點完整", "學習者理解與練習題"),
        ("主要難點", "事實正確性", "教學設計")]
tx, ty = 0.85, 2.05
colw = [1.20, 2.65, 2.30]
rrect(s, tx, ty, sum(colw), 0.44, fill=INK, shape=MSO_SHAPE.RECTANGLE)
heads = ["", "論文整理（本項目）", "Lecture note（不做）"]
cx = tx
for i, hd in enumerate(heads):
    tf = tbox(s, cx + 0.16, ty + 0.11, colw[i] - 0.2, 0.3)
    add(tf, hd, size=9.5, bold=True, color=WHITE, first=True)
    cx += colw[i]
for r, row in enumerate(rows):
    y = ty + 0.44 + r * 0.74
    bg = PANEL if r % 2 == 0 else WHITE
    rrect(s, tx, y, sum(colw), 0.74, fill=bg, line=RGBColor(0xE2, 0xE8, 0xEE), lw=0.6, shape=MSO_SHAPE.RECTANGLE)
    cx = tx
    for i, cell in enumerate(row):
        tf = tbox(s, cx + 0.16, y + 0.10, colw[i] - 0.26, 0.6)
        if i == 0:
            add(tf, cell, size=9.5, bold=True, color=GREY, first=True)
        elif i == 1:
            add(tf, cell, size=9.5, color=INK, line=1.25, first=True)
        else:
            add(tf, cell, size=9.5, color=LGREY, line=1.25, first=True)
        cx += colw[i]

# RQ
rrect(s, 7.32, 2.05, 5.16, 3.45, fill=PANEL2)
tf = tbox(s, 7.58, 2.22, 4.6, 0.3)
add(tf, "研究問題", size=10, bold=True, color=GOLD, first=True)
rqs = [("RQ1", "如何為每條 bullet 建立可驗證的來源錨點？"),
       ("RQ2", "如何把版面可用性變成可計算的指標？"),
       ("RQ3", "檢查—修復迴圈能否同時提升兩者，而成本可接受？")]
for i, (k, q) in enumerate(rqs):
    y = 2.60 + i * 0.98
    tf = tbox(s, 7.58, y, 0.6, 0.3)
    add(tf, k, size=11, bold=True, color=RED, first=True)
    tf = tbox(s, 8.16, y - 0.02, 4.05, 0.9)
    add(tf, q, size=10.5, color=INK2, line=1.35, first=True)

footer(s, 10)

# =====================================================================
# 11 CONTRIBUTIONS
# =====================================================================
s = new_slide()
header(s, "03 · 研究定位", "三個貢獻")

contribs = [
    ("01", "可計算的「可信度」框架",
     "把「內容可信」與「版面可用」各變成一個可自動量度的比率：\n內容可追溯率（traceable bullet ratio）＋ 版面可用率（deck-ready rate）。",
     "交付：指標定義 + 開源計算套件"),
    ("02", "逐步檢查—修復迴圈",
     "每次只改一個元素、改完立即渲染驗證，令錯誤可歸因；\n並報告迭代的收斂行為，而不是只報最終分數。",
     "交付：agent 實作 + 收斂曲線"),
    ("03", "公開資料集與評測協定",
     "150 版人手標註的違規 ground truth、30 人 pairwise 偏好資料、\n以及可重跑的評測腳本。",
     "交付：dataset + benchmark script"),
]
for i, (num, title, body, deliver) in enumerate(contribs):
    y = 1.85 + i * 1.62
    rrect(s, 0.85, y, 11.6, 1.42, fill=PANEL if i % 2 == 0 else WHITE,
          line=RGBColor(0xE2, 0xE8, 0xEE), lw=0.7)
    tf = tbox(s, 1.10, y + 0.20, 0.9, 0.6)
    add(tf, num, size=22, bold=True, color=GOLD, first=True)
    tf = tbox(s, 2.00, y + 0.18, 4.1, 0.4)
    add(tf, title, size=14, bold=True, color=INK, first=True)
    tf = tbox(s, 2.00, y + 0.62, 6.6, 0.7)
    add(tf, body, size=9.8, color=INK2, line=1.35, first=True)
    tf = tbox(s, 8.95, y + 0.55, 3.3, 0.55)
    add(tf, deliver, size=9, bold=True, color=TEAL, line=1.3, first=True)

footer(s, 11)

# =====================================================================
# 12 PIPELINE
# =====================================================================
s = new_slide()
header(s, "04 · 方法", "系統架構", "金色 = 本項目的新貢獻；其餘沿用已被驗證的成熟做法")

stages = [("PDF", "輸入", False), ("解析", "section / figure / table", False),
          ("證據圖", "claim–evidence", True), ("敘事規劃", "篇幅分配", True),
          ("版面約束", "容量感知", True), ("渲染", "python-pptx", False),
          ("檢查", "違規偵測", True), ("修復", "局部重寫", True)]
bw2, gap3 = 1.36, 0.15
x0 = 0.85
for i, (name, desc, new) in enumerate(stages):
    x = x0 + i * (bw2 + gap3)
    fill = GOLDBG if new else PANEL2
    line = GOLD if new else RGBColor(0xD8, 0xE0, 0xE8)
    rrect(s, x, 2.05, bw2, 1.32, fill=fill, line=line, lw=1.2 if new else 0.8)
    tf = tbox(s, x + 0.08, 2.22, bw2 - 0.16, 0.42)
    add(tf, name, size=12, bold=True, color=INK, align=PP_ALIGN.CENTER, first=True)
    tf = tbox(s, x + 0.08, 2.64, bw2 - 0.16, 0.6)
    add(tf, desc, size=7.5, color=GREY, align=PP_ALIGN.CENTER, line=1.2, first=True)
    if i < len(stages) - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw2 + 0.015), Inches(2.60),
                                Inches(0.12), Inches(0.20))
        ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor(0xC5, 0xCF, 0xD9)
        ar.line.fill.background(); ar.shadow.inherit = False; nostyle(ar)

principles = [("限制輸出空間", "模型只寫結構化內容，座標與樣式由模板負責"),
              ("一次只改一個元素", "局部重寫 + 即時驗證，錯誤才可歸因"),
              ("所有判斷都要有數字", "不用「看起來不錯」，只用可量度的違規與比率")]
for i, (t, d) in enumerate(principles):
    x = 0.85 + i * 3.92
    rrect(s, x, 3.95, 3.66, 1.15, fill=WHITE, line=RGBColor(0xDF, 0xE5, 0xEB), lw=0.8)
    tf = tbox(s, x + 0.26, 4.12, 3.1, 0.34)
    add(tf, t, size=11.5, bold=True, color=INK, first=True)
    tf = tbox(s, x + 0.26, 4.50, 3.1, 0.95)
    add(tf, d, size=9.8, color=GREY, line=1.35, first=True)

tf = tbox(s, 0.85, 5.45, 11.6, 0.7)
add(tf, "輸出的不是圖片，而是可編輯的 .pptx：版面由約束求解決定，內容由證據圖決定。",
    size=11, color=INK2, first=True)
footer(s, 12)

# =====================================================================
# 13 EVALUATION
# =====================================================================
s = new_slide()
header(s, "04 · 方法", "如何證明有效", "三軌並行，避免「用 LLM 評 LLM」的循環論證")

tracks = [("自動（幾何）", "違規數、版面可用率", "8 篇 held-out、150 版標註", "完全確定性、可重跑"),
          ("任務式（問答）", "內容傳達正確率", "每份 deck 10 題 × 8 篇", "測「觀眾能否從成品得到正確資訊」"),
          ("人評", "pairwise 偏好、可講性", "30 人", "設計沒有 ground truth，只能相對比較")]
hdrs = ["軌道", "指標", "樣本", "為何需要"]
colx = [0.85, 3.55, 6.75, 9.55]
colw3 = [2.60, 3.10, 2.70, 2.90]
rrect(s, 0.85, 2.05, 11.6, 0.42, fill=INK, shape=MSO_SHAPE.RECTANGLE)
for i, h in enumerate(hdrs):
    tf = tbox(s, colx[i] + 0.12, 2.13, colw3[i] - 0.2, 0.3)
    add(tf, h, size=9, bold=True, color=WHITE, first=True)
for r, row in enumerate(tracks):
    y = 2.47 + r * 0.66
    rrect(s, 0.85, y, 11.6, 0.66, fill=PANEL if r % 2 == 0 else WHITE,
          line=RGBColor(0xE4, 0xEA, 0xF0), lw=0.6, shape=MSO_SHAPE.RECTANGLE)
    for i, cell in enumerate(row):
        tf = tbox(s, colx[i] + 0.12, y + 0.17, colw3[i] - 0.22, 0.4)
        add(tf, cell, size=8.8 if i == 3 else 9.5,
            bold=(i == 0), color=INK if i < 3 else GREY, line=1.25, first=True)

rrect(s, 0.85, 4.95, 5.7, 1.12, fill=PANEL2)
tf = tbox(s, 1.08, 5.09, 5.2, 0.25)
add(tf, "BASELINES", size=8.5, bold=True, color=GOLD, first=True)
tf = tbox(s, 1.08, 5.35, 5.2, 0.85)
add(tf, "① Naive single-pass（一次生成，無檢查）\n② PPTAgent（領域最強可複現系統）", size=9.5, color=INK2, line=1.4, first=True)

rrect(s, 6.78, 4.95, 5.67, 1.12, fill=PANEL2)
tf = tbox(s, 7.01, 5.09, 5.2, 0.25)
add(tf, "ABLATION", size=8.5, bold=True, color=GOLD, first=True)
tf = tbox(s, 7.01, 5.35, 5.2, 0.85)
add(tf, "移除引用錨點 / 移除檢查 / 移除修復\n→ 逐項歸因每個模組的貢獻", size=9.5, color=INK2, line=1.4, first=True)
footer(s, 13)

# =====================================================================
# 14 TIMELINE
# =====================================================================
s = new_slide()
header(s, "05 · 計劃", "20 週時間表")

phases = [("文獻與範圍", 0, 2, "W1–2"),
          ("環境與 spike", 2, 2, "W3–4"),
          ("解析與證據圖", 4, 3, "W5–7"),
          ("版面與渲染", 7, 4, "W8–11"),
          ("檢查與修復", 11, 3, "W12–14"),
          ("實驗與人評", 14, 4, "W15–18"),
          ("報告與簡報", 18, 2, "W19–20")]
gx, gw = 2.75, 9.4
for wk in range(0, 21, 5):
    x = gx + gw * wk / 20
    ln = s.shapes.add_connector(1, Inches(x), Inches(1.85), Inches(x), Inches(5.95))
    ln.line.color.rgb = RGBColor(0xEC, 0xF0, 0xF4); ln.line.width = Pt(0.75)
    tf = tbox(s, x - 0.35, 1.60, 0.7, 0.22)
    add(tf, f"W{wk}" if wk else "", size=7.5, color=LGREY, align=PP_ALIGN.CENTER, first=True)

for i, (name, start, dur, lab) in enumerate(phases):
    y = 2.00 + i * 0.56
    x = gx + gw * start / 20
    w = max(gw * dur / 20, 0.9)
    col = GOLD if i in (2, 4) else TEAL
    rrect(s, x, y, w, 0.36, fill=col, line=None, radius=0.25)
    tf = tbox(s, x + 0.12, y + 0.075, w - 0.15, 0.25)
    add(tf, lab, size=7.5, bold=True, color=WHITE, first=True)
    tf = tbox(s, 0.85, y + 0.05, 1.85, 0.3)
    add(tf, name, size=9.5, bold=True, color=INK, first=True)

mss = [("2026-11-05", "Proposal 提交"), ("2027-01-07", "Interim Report"),
       ("2027-04-08", "Final Report + Turnitin + AI Report")]
for i, (dt, nm) in enumerate(mss):
    x = 0.85 + i * 3.92
    rrect(s, x, 6.00, 3.66, 0.76, fill=PANEL)
    rrect(s, x, 6.00, 0.05, 0.76, fill=RED, shape=MSO_SHAPE.RECTANGLE)
    tf = tbox(s, x + 0.22, 6.08, 1.6, 0.3)
    add(tf, dt, size=9.5, bold=True, color=RED, first=True)
    tf = tbox(s, x + 0.22, 6.40, 3.3, 0.3)
    add(tf, nm, size=8.5, color=GREY, first=True)
footer(s, 14)

# =====================================================================
# 15 REFERENCES (APA 7)
# =====================================================================
s = new_slide()
header(s, "附錄 · APA 7th", "引用文獻")

refs = [
 "Chen, H., Zhou, X., Zhou, W., Shao, X., Zhu, Y., Wang, B., Hong, J., Jia, A., & Wu, F. (2026). X+Slides: Benchmarking audience-conditioned slide generation (arXiv:2606.19256). arXiv. https://doi.org/10.48550/arXiv.2606.19256",
 "Fu, T.-J., Wang, W. Y., McDuff, D., & Song, Y. (2021). DOC2PPT: Automatic presentation slides generation from scientific documents (arXiv:2101.11796). arXiv. https://doi.org/10.48550/arXiv.2101.11796",
 "Ge, J., Wang, Z. Z., Zhou, X., Peng, Y.-H., Subramanian, S., Tan, Q., Sap, M., Suhr, A., Fried, D., Neubig, G., & Darrell, T. (2025). AutoPresent: Designing structured visuals from scratch (arXiv:2501.00912). arXiv. https://doi.org/10.48550/arXiv.2501.00912",
 "Hu, H., Dang, C., Liu, Y., Kang, H., He, C., & Li, W. (2026). PosterMELD: Multi-agent paper-to-poster generation for controllable design diversity with editable print-ready outputs (arXiv:2608.02218). arXiv. https://doi.org/10.48550/arXiv.2608.02218",
 "Huang, Z., Liu, X., Hu, T., Zhang, K., & Liu, Y. (2025). PPTBench: Towards holistic evaluation of large language models for PowerPoint layout and design understanding (arXiv:2512.02624). arXiv. https://doi.org/10.48550/arXiv.2512.02624",
 "Jung, K., Cho, H., Yun, J., Yang, S., Jang, J., & Choo, J. (2025). Talk to your slides: High-efficiency slide editing via language-driven structured data manipulation (arXiv:2505.11604). arXiv. https://doi.org/10.48550/arXiv.2505.11604",
 "Liang, X., Zhang, Z., Zhang, X., Su, H., Xu, Y., Sun, S., & You, C. (2025). SlideGen: Collaborative multimodal agents for scientific slide generation (arXiv:2512.04529). arXiv. https://doi.org/10.48550/arXiv.2512.04529",
 "Ozden, T. C., VS, S., Horoz, F., Kara, O., Kim, J., & Rehg, J. M. (2026a). Narrative-driven paper-to-slide generation via ArcDeck (arXiv:2604.11969). arXiv. https://doi.org/10.48550/arXiv.2604.11969",
 "Ozden, T. C., VS, S., Horoz, F., Kara, O., Hakkani-Tür, D., Kim, J., & Rehg, J. M. (2026b). ConvDeck: Conversational paper-to-slide generation via stage-specific user feedback (arXiv:2609.00226). arXiv. https://doi.org/10.48550/arXiv.2609.00226",
 "Pang, W., Lin, K. Q., Jian, X., He, X., & Torr, P. (2025). Paper2Poster: Towards multimodal poster automation from scientific papers (arXiv:2505.21497). arXiv. https://doi.org/10.48550/arXiv.2505.21497",
 "Sun, E., Hou, Y., Wang, D., Zhang, Y., & Wang, N. X. R. (2021). D2S: Document-to-slide generation via query-based text summarization (arXiv:2105.03664). arXiv. https://doi.org/10.48550/arXiv.2105.03664",
 "Tian, M., Zeng, Z., Yang, Y., Gao, X., Li, Y., Huang, Z., Wang, X., Lv, C., Cheng, M., Liu, B., Qiu, K., Dai, Q., Chen, D., Dong, Y., Zheng, X., Li, J., & Luo, C. (2026). ReDeck: Step-level render-grounded refinement for document-to-slide generation (arXiv:2609.00194). arXiv. https://doi.org/10.48550/arXiv.2609.00194",
 "Vats, V., Sharma, K., & Wang, Y. (2026). SlideLab: Audience-centered scientific slide generation and evaluation (arXiv:2609.30294). arXiv. https://doi.org/10.48550/arXiv.2609.30294",
 "Xie, E., Waterfield, D., Kennedy, M., & Zhang, A. (2025). SlideBot: A multi-agent framework for generating informative, reliable, multi-modal presentations (arXiv:2511.09804). arXiv. https://doi.org/10.48550/arXiv.2511.09804",
 "Yang, M., Zhang, Z., Li, J., Liu, H., Cai, Y., & Zheng, W. (2026). DeepSlide: From artifacts to presentation delivery (arXiv:2605.15202). arXiv. https://doi.org/10.48550/arXiv.2605.15202",
 "Zheng, H., Guan, X., Kong, H., Zheng, J., Zhou, W., Lin, H., Lu, Y., He, B., Han, X., & Sun, L. (2025). PPTAgent: Generating and evaluating presentations beyond text-to-slides (arXiv:2501.03936). arXiv. https://doi.org/10.48550/arXiv.2501.03936",
 "Zheng, H., Mo, G., Yan, X., Yuan, Q., Zhang, W., Chen, X., Lu, Y., Lin, H., Han, X., & Sun, L. (2026a). DeepPresenter: Environment-grounded reflection for agentic presentation generation (arXiv:2602.22839). arXiv. https://doi.org/10.48550/arXiv.2602.22839",
 "Zheng, H., Wang, F., Xiong, T., Yu, Y., Qian, S., Yu, H., Schwing, A., Nahrstedt, K., & Wu, M. (2026b). SLIDEFORGE: An LLM agent for controllable editing of slides as structured artifacts (arXiv:2609.03109). arXiv. https://doi.org/10.48550/arXiv.2609.03109",
]
half = 9
cols = [refs[:half], refs[half:]]
for c, group in enumerate(cols):
    x = 0.85 + c * 5.95
    tf = tbox(s, x, 1.72, 5.55, 5.1)
    for i, r in enumerate(group):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0 if i == 0 else 4)
        p.line_spacing = 1.06
        run = p.add_run(); run.text = r
        run.font.name = FONT; run.font.size = Pt(7.6); run.font.color.rgb = INK2
footer(s, 15)

out = "/Users/yeetungchan/Documents/deepseek-harness/default-workspace/capstone-ppt-agent/deck/proposal_deck.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
