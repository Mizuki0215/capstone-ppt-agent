"""Projection of the deck IR onto a real .pptx (python-pptx, no clip art)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

from . import spec
from .spec import Block, Deck, Slide, para_specs


def _rgb(c) -> RGBColor:
    return RGBColor(int(c[0]), int(c[1]), int(c[2]))


def _bullet_xml(mar_l: int, indent: int, gap_pt: float, first: bool) -> str:
    space = "" if first else f'<a:spcBef><a:spcPts val="{int(gap_pt * 100)}"/></a:spcBef>'
    return (
        f'<a:pPr {nsdecls("a")} marL="{mar_l}" indent="{indent}">{space}'
        f'<a:buFont typeface="Arial"/><a:buChar char="\u2022"/></a:pPr>'
    )


def _apply_bullet(p, first: bool) -> None:
    mar_l = int(Inches(spec.BULLET_INDENT))
    xml = _bullet_xml(mar_l, -mar_l, spec.PARA_GAP_PT, first)
    pPr = p._p.find(qn("a:pPr"))
    if pPr is not None:
        p._p.remove(pPr)
    p._p.insert(0, parse_xml(xml))


def _write_block(tf, block: Block) -> None:
    tf.word_wrap = True
    tf.margin_left = Inches(spec.BOX_MX)
    tf.margin_right = Inches(spec.BOX_MX)
    tf.margin_top = Inches(spec.BOX_MY)
    tf.margin_bottom = Inches(spec.BOX_MY)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, (text, size, bold, color) in enumerate(para_specs(block)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.name = "Arial"
        f.color.rgb = _rgb(color)
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT}[block.align]
        if block.kind == "bullets":
            _apply_bullet(p, i == 0)
        elif i > 0:
            p.space_before = Pt(spec.PARA_GAP_PT)


def _add_block(slide, block: Block, name: str) -> None:
    tb = slide.shapes.add_textbox(Inches(block.x), Inches(block.y),
                                  Inches(block.w), Inches(block.h))
    tb.name = name
    if block.fill is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = _rgb(block.fill)
        tb.line.fill.background()
    _write_block(tb.text_frame, block)


def _rule(slide, x, y, w, h, color, name):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.name = name
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False


def _render(slide, s: Slide, idx: int) -> None:
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb(spec.PAPER)

    if s.layout == "cover":
        _rule(slide, 0.92, 1.86, 0.62, 0.09, spec.ACCENT, "cover_rule")
        tb = slide.shapes.add_textbox(Inches(0.9), Inches(2.05),
                                      Inches(11.6), Inches(1.9))
        tb.name = f"t{idx}"
        _write_paragraphs(tb.text_frame, [(s.title, 44.0, True, spec.NAVY)],
                          gap_pt=0)
        for j, b in enumerate(s.blocks):
            _add_block(slide, b, f"b{idx}_{j}")
        return

    _rule(slide, 0, 0, spec.SLIDE_W, 0.10, spec.NAVY, "bar")
    tb = slide.shapes.add_textbox(Inches(spec.MARGIN), Inches(0.40),
                                  Inches(spec.TEXT_W), Inches(0.86))
    tb.name = f"t{idx}"
    _write_paragraphs(tb.text_frame, [(s.title, 30.0, True, spec.INK)], gap_pt=0)
    _rule(slide, spec.MARGIN, 1.33, 1.55, 0.035, spec.ACCENT, "rule")
    for j, b in enumerate(s.blocks):
        _add_block(slide, b, f"b{idx}_{j}")

    pn = slide.shapes.add_textbox(Inches(spec.SLIDE_W - 1.1), Inches(7.02),
                                  Inches(0.6), Inches(0.3))
    pn.name = f"f{idx}"
    _write_paragraphs(pn.text_frame, [(f"{idx + 1:02d}", 11.0, False, spec.MUTED)],
                      gap_pt=0, align=PP_ALIGN.RIGHT)


def _write_paragraphs(tf, paras, gap_pt, align=PP_ALIGN.LEFT) -> None:
    tf.word_wrap = True
    tf.auto_size = None
    for i, (text, size, bold, color) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Arial"
        run.font.color.rgb = _rgb(color)
        p.alignment = align
        if i > 0 and gap_pt:
            p.space_before = Pt(gap_pt)


def build(deck: Deck, path: str | Path) -> Path:
    prs = Presentation()
    prs.slide_width = Emu(int(spec.SLIDE_W * 914400))
    prs.slide_height = Emu(int(spec.SLIDE_H * 914400))
    blank = prs.slide_layouts[6]
    for i, s in enumerate(deck.slides):
        slide = prs.slides.add_slide(blank)
        _render(slide, s, i)
        if s.notes:
            slide.notes_slide.notes_text_frame.text = "\n".join(s.notes)
    prs.core_properties.title = deck.title
    prs.core_properties.author = "SRR prototype"
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
