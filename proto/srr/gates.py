"""G1–G6: the deterministic predicates. No model, no API, bit-for-bit rerunnable.

Each gate reads the *native* .pptx object model (exact pt sizes, bounding
boxes, solid fills), not a rendered image, so a measurement is a property of
the file rather than of the renderer.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_FILL

from . import metrics, spec

NAME_RE = re.compile(r"^([tb])(\d+)(?:_(\d+))?$")

GATE_NAMES = {
    "G1": "Containment",
    "G2": "Overlap",
    "G3": "Font size",
    "G4": "Contrast",
    "G5": "Text fit",
    "G6": "Density",
}
EMU_IN = 914400.0
EPS = 1e-6


@dataclass
class Para:
    text: str
    size: float
    bold: bool
    color: tuple
    indent: float
    gap_pt: float


@dataclass
class TShape:
    name: str
    is_title: bool
    block_idx: int
    x: float
    y: float
    w: float
    h: float
    paras: list[Para]
    fill: tuple | None

    @property
    def right(self) -> float:
        return self.x + self.w

    @property
    def bottom(self) -> float:
        return self.y + self.h

    @property
    def area(self) -> float:
        return self.w * self.h


@dataclass
class Violation:
    gate: str
    slide: int
    shapes: list[str]
    value: float
    threshold: float
    detail: str


@dataclass
class SlideResult:
    index: int
    title: str
    gates: dict[str, bool]
    violations: list[Violation] = field(default_factory=list)
    stats: dict = field(default_factory=dict)

    @property
    def passed(self) -> bool:
        return all(self.gates.values())


@dataclass
class DeckResult:
    slides: list[SlideResult]

    @property
    def n(self) -> int:
        return len(self.slides)

    @property
    def n_ready(self) -> int:
        return sum(1 for s in self.slides if s.passed)

    @property
    def srr(self) -> float:
        return self.n_ready / self.n if self.n else 0.0

    @property
    def deck_ready(self) -> bool:
        return self.n > 0 and self.n_ready == self.n

    def per_gate(self) -> dict[str, float]:
        out = {}
        for g in GATE_NAMES:
            n = len(self.slides)
            out[g] = sum(1 for s in self.slides if s.gates.get(g, True)) / n if n else 0.0
        return out

    def gate_counts(self) -> dict[str, int]:
        counts = {g: 0 for g in GATE_NAMES}
        for s in self.slides:
            for v in s.violations:
                counts[v.gate] += 1
        return counts

    def as_dict(self) -> dict:
        return {
            "srr": self.srr,
            "deck_ready": self.deck_ready,
            "n_slides": self.n,
            "n_ready": self.n_ready,
            "per_gate": self.per_gate(),
            "violations": self.gate_counts(),
            "slides": [
                {"index": s.index, "title": s.title, "passed": s.passed,
                 "gates": s.gates, "stats": s.stats,
                 "violations": [v.__dict__ for v in s.violations]}
                for s in self.slides
            ],
        }


def _solid_fill(shape) -> tuple | None:
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            c = shape.fill.fore_color.rgb
            return (c[0], c[1], c[2])
    except (AttributeError, TypeError, ValueError, KeyError):
        pass
    return None


def _shape(shape, idx: int) -> TShape | None:
    if not getattr(shape, "has_text_frame", False) or not shape.text_frame.text.strip():
        return None
    m = NAME_RE.match(shape.name or "")
    if not m:
        return None
    kind, sidx, b = m.group(1), int(m.group(2)), m.group(3)
    if sidx != idx:
        return None
    paras: list[Para] = []
    for p in shape.text_frame.paragraphs:
        text = "".join(r.text for r in p.runs)
        if not text.strip():
            continue
        size, bold, color = 20.0, False, spec.INK
        for r in p.runs:
            if r.font.size is not None:
                size = float(r.font.size.pt)
            if r.font.bold is not None:
                bold = bool(r.font.bold)
            try:
                c = r.font.color
                if c is not None and c.type is not None and c.rgb is not None:
                    color = (c.rgb[0], c.rgb[1], c.rgb[2])
            except (AttributeError, TypeError, ValueError):
                pass
            break
        indent = 0.0
        gap = 0.0
        pPr = p._p.pPr
        if pPr is not None:
            if pPr.get("marL"):
                indent = int(pPr.get("marL")) / EMU_IN
            spc = pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spcBef")
            if spc is not None:
                pts = spc.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}spcPts")
                if pts is not None:
                    gap = int(pts.get("val")) / 100.0
        paras.append(Para(text, size, bold, color, indent, gap))
    if not paras:
        return None
    return TShape(
        name=shape.name, is_title=(kind == "t"), block_idx=int(b) if b is not None else -1,
        x=shape.left / EMU_IN, y=shape.top / EMU_IN,
        w=shape.width / EMU_IN, h=shape.height / EMU_IN,
        paras=paras, fill=_solid_fill(shape),
    )


def _g5_ratio(sh: TShape) -> float:
    indent = sh.paras[0].indent if sh.paras else 0.0
    gap = max((p.gap_pt for p in sh.paras), default=0.0)
    need = metrics.required_height([(p.text, p.size, p.bold) for p in sh.paras],
                                   sh.w, indent, gap)
    return need / metrics.usable_height(sh.h)


def measure_prs(prs) -> DeckResult:
    results: list[SlideResult] = []
    for i, slide in enumerate(prs.slides):
        shapes = [s for s in (_shape(sh, i) for sh in slide.shapes) if s]
        title = next((s.paras[0].text for s in shapes if s.is_title), "")
        v: list[Violation] = []
        gates = {g: True for g in GATE_NAMES}

        # G1 containment
        for sh in shapes:
            over = max(-sh.x, -sh.y, sh.right - spec.SLIDE_W, sh.bottom - spec.SLIDE_H)
            if over > EPS:
                gates["G1"] = False
                v.append(Violation("G1", i, [sh.name], round(over, 3), 0.0,
                                   f"bbox exceeds slide by {over:.2f} in"))

        # G2 overlap (text vs text)
        for a in range(len(shapes)):
            for b in range(a + 1, len(shapes)):
                s1, s2 = shapes[a], shapes[b]
                iw = min(s1.right, s2.right) - max(s1.x, s2.x)
                ih = min(s1.bottom, s2.bottom) - max(s1.y, s2.y)
                if iw <= 0 or ih <= 0:
                    continue
                frac = (iw * ih) / max(EPS, min(s1.area, s2.area))
                if frac > spec.OVERLAP_FRAC:
                    gates["G2"] = False
                    v.append(Violation("G2", i, [s1.name, s2.name], round(frac, 3),
                                       spec.OVERLAP_FRAC,
                                       f"{frac:.1%} of the smaller box overlapped"))

        # G3 minimum type size
        body = [s for s in shapes if not s.is_title]
        mins = [(p.size, s.name) for s in body for p in s.paras]
        if mins:
            smallest, who = min(mins)
            if smallest < spec.MIN_BODY_PT - EPS:
                gates["G3"] = False
                v.append(Violation("G3", i, [who], smallest, spec.MIN_BODY_PT,
                                   f"body text at {smallest:g} pt"))
        for s in shapes:
            if s.is_title and s.paras[0].size < spec.MIN_TITLE_PT - EPS:
                gates["G3"] = False
                v.append(Violation("G3", i, [s.name], s.paras[0].size,
                                   spec.MIN_TITLE_PT,
                                   f"title at {s.paras[0].size:g} pt"))

        # G4 contrast
        for s in shapes:
            bg = s.fill or spec.PAPER
            worst, bad = 99.0, False
            for p in s.paras:
                r = metrics.contrast(p.color, bg)
                worst = min(worst, r)
                if r < metrics.threshold_for(p.size, p.bold) - EPS:
                    bad = True
            if bad:
                gates["G4"] = False
                v.append(Violation("G4", i, [s.name], round(worst, 2),
                                   spec.CONTRAST_AA, f"contrast {worst:.2f}:1"))

        # G5 rendered text fit
        ratios = []
        for s in shapes:
            r = _g5_ratio(s)
            ratios.append(r)
            if r > spec.FIT_RATIO + EPS:
                gates["G5"] = False
                v.append(Violation("G5", i, [s.name], round(r, 3), spec.FIT_RATIO,
                                   f"needs {r:.0%} of the box height"))
        # G6 density
        chars = sum(len(p.text) for s in body for p in s.paras)
        max_bul = max((len(s.paras) for s in body), default=0)
        if chars > spec.MAX_CHARS:
            gates["G6"] = False
            v.append(Violation("G6", i, [s.name for s in body], chars,
                               spec.MAX_CHARS, f"{chars} chars on the slide"))
        elif max_bul > spec.MAX_BULLETS:
            gates["G6"] = False
            v.append(Violation("G6", i, [x.name for x in body], max_bul,
                               spec.MAX_BULLETS, f"{max_bul} bullets in one block"))

        results.append(SlideResult(
            index=i, title=title, gates=gates, violations=v,
            stats={"fit_max": round(max(ratios), 3) if ratios else 0.0,
                   "chars": chars, "bullets": max_bul,
                   "min_pt": min((p.size for s in body for p in s.paras), default=0)}))
    return DeckResult(results)


def measure(path: str | Path) -> DeckResult:
    return measure_prs(Presentation(str(path)))
