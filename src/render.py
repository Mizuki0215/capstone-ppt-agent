#!/usr/bin/env python3
"""Stage 3 - render a slide plan into an editable .pptx with speaker notes.

Runs on a Python that has python-pptx.
"""
import argparse
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x10, 0x23, 0x3A)
GREY = RGBColor(0x5A, 0x6B, 0x7C)
GOLD = RGBColor(0xC0, 0x8A, 0x2E)
PANEL = RGBColor(0xF4, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Helvetica Neue"


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, attr, 0)
    return tf


def para(tf, text, size, bold=False, color=INK, first=False, space_before=0, line=1.2,
         align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.line_spacing = line
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def body_size(n_bullets, n_chars):
    if n_chars > 520:
        return 15
    if n_chars > 420:
        return 17
    if n_bullets >= 6:
        return 18
    return 20


def build(plan, out):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---- title slide
    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.3), Inches(7.5))
    band.fill.solid(); band.fill.fore_color.rgb = INK
    band.line.fill.background()
    acc = s.shapes.add_shape(1, Inches(0.3), Inches(0), Inches(0.06), Inches(7.5))
    acc.fill.solid(); acc.fill.fore_color.rgb = GOLD
    acc.line.fill.background()

    tf = textbox(s, 1.0, 1.6, 10.6, 3.0)
    para(tf, plan["title"], 34, bold=True, first=True, line=1.15)
    tf = textbox(s, 1.0, 4.9, 10.0, 0.5)
    para(tf, "A generated deck - draft, not final", 14, color=GREY, first=True)
    add_notes(s, "Title slide. State the paper, then the one-sentence claim of the talk.")

    # ---- content slides
    for i, sl in enumerate(plan["slides"], start=2):
        s = prs.slides.add_slide(blank)
        tf = textbox(s, 0.9, 0.55, 11.5, 0.9)
        para(tf, sl["title"], 28, bold=True, first=True, line=1.1)

        rule = s.shapes.add_shape(1, Inches(0.9), Inches(1.42), Inches(11.5), Inches(0.02))
        rule.fill.solid(); rule.fill.fore_color.rgb = RGBColor(0xDD, 0xE2, 0xE8)
        rule.line.fill.background()

        bullets = sl["bullets"]
        total = sum(len(b) for b in bullets)
        size = body_size(len(bullets), total)
        tf = textbox(s, 0.95, 1.75, 11.4, 4.5)
        for k, b in enumerate(bullets):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_before = Pt(0 if k == 0 else 12)
            p.line_spacing = 1.25
            mark = p.add_run()
            mark.text = "-  "
            mark.font.name = FONT
            mark.font.size = Pt(size)
            mark.font.bold = True
            mark.font.color.rgb = GOLD
            run = p.add_run()
            run.text = b
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = INK

        if sl.get("figure"):
            panel = s.shapes.add_shape(1, Inches(0.95), Inches(6.15), Inches(11.4), Inches(0.75))
            panel.fill.solid(); panel.fill.fore_color.rgb = PANEL
            panel.line.fill.background()
            tf = textbox(s, 1.15, 6.28, 11.0, 0.5)
            para(tf, sl["figure"], 11, color=GREY, first=True, line=1.2)

        tf = textbox(s, 11.9, 7.05, 0.6, 0.3)
        para(tf, f"{i:02d}", 9, color=GREY, first=True, align=PP_ALIGN.RIGHT)
        add_notes(s, sl.get("notes", ""))

    prs.save(out)
    return len(prs.slides._sldIdLst)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("plan_json")
    ap.add_argument("-o", "--out", required=True)
    args = ap.parse_args()
    plan = json.load(open(args.plan_json))
    n = build(plan, args.out)
    print(f"slides={n} -> {args.out}")


if __name__ == "__main__":
    main()
