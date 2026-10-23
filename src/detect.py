#!/usr/bin/env python3
"""Detector v1 - exact checks straight off the .pptx object model.

Implements the four gates that need no font metrics and no rendering:
  G1 containment, G2 overlap, G3 font size, G6 density.
G4 (contrast) and G5 (text fit) need colour resolution and font metrics; not yet.

Usage:
  python detect.py deck.pptx [--json out.json]
"""
import argparse
import json
from itertools import combinations

from pptx import Presentation
from pptx.util import Emu

BODY_MIN_PT = 18
MAX_CHARS = 140
MAX_BULLETS = 6
OVERLAP_FRAC = 0.05


def emu_box(shape):
    return (shape.left or 0, shape.top or 0,
            (shape.left or 0) + (shape.width or 0),
            (shape.top or 0) + (shape.height or 0))


def area(b):
    return max(0, b[2] - b[0]) * max(0, b[3] - b[1])


def intersection(a, b):
    x1, y1 = max(a[0], b[0]), max(a[1], b[1])
    x2, y2 = min(a[2], b[2]), min(a[3], b[3])
    if x2 <= x1 or y2 <= y1:
        return 0
    return (x2 - x1) * (y2 - y1)


def text_shapes(slide):
    out = []
    for sh in slide.shapes:
        if not getattr(sh, "has_text_frame", False):
            continue
        if not sh.text_frame.text.strip():
            continue
        out.append(sh)
    return out


def min_font_pt(shape):
    sizes = []
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.text.strip() and r.font.size is not None:
                sizes.append(r.font.size.pt)
    return min(sizes) if sizes else None


def check_slide(slide, sw, sh):
    v = []
    texts = text_shapes(slide)

    # G1 containment
    for s in slide.shapes:
        b = emu_box(s)
        if b[0] < 0 or b[1] < 0 or b[2] > sw or b[3] > sh:
            v.append({"gate": "G1", "detail": f"shape outside slide: {s.name}"})

    # G2 overlap (text vs text)
    for a, b in combinations(texts, 2):
        ba, bb = emu_box(a), emu_box(b)
        inter = intersection(ba, bb)
        if inter <= 0:
            continue
        smaller = min(area(ba), area(bb))
        if smaller and inter / smaller > OVERLAP_FRAC:
            v.append({"gate": "G2",
                      "detail": f"{a.name} overlaps {b.name} ({inter/smaller:.0%})"})

    # G3 font size + G6 density
    total_chars = 0
    total_bullets = 0
    for s in texts:
        pt = min_font_pt(s)
        if pt is not None and pt < BODY_MIN_PT:
            v.append({"gate": "G3", "detail": f"{s.name} at {pt:g} pt"})
        total_chars += len(s.text_frame.text)
        for p in s.text_frame.paragraphs:
            if p.text.strip():
                total_bullets += 1

    if total_chars > MAX_CHARS:
        v.append({"gate": "G6", "detail": f"{total_chars} chars"})
    if total_bullets > MAX_BULLETS:
        v.append({"gate": "G6", "detail": f"{total_bullets} bullets"})

    return v, total_chars


def analyse(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    rows = []
    for i, slide in enumerate(prs.slides, start=1):
        v, chars = check_slide(slide, sw, sh)
        rows.append({"slide": i, "violations": v, "chars": chars,
                     "clean": len(v) == 0})
    by_gate = {}
    for r in rows:
        for g in {x["gate"] for x in r["violations"]}:
            by_gate[g] = by_gate.get(g, 0) + 1
    clean = sum(1 for r in rows if r["clean"])
    return {
        "deck": path,
        "slides": len(rows),
        "clean_slides": clean,
        "srr": round(clean / len(rows), 3) if rows else 0,
        "violations_total": sum(len(r["violations"]) for r in rows),
        "slides_with_violation_by_gate": by_gate,
        "detail": rows,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx", nargs="+")
    ap.add_argument("--json")
    args = ap.parse_args()
    results = [analyse(p) for p in args.pptx]
    for r in results:
        print(f"\n{r['deck']}")
        print(f"  slides={r['slides']}  clean={r['clean_slides']}  SRR={r['srr']}")
        print(f"  violations={r['violations_total']}  by gate={r['slides_with_violation_by_gate']}")
        for row in r["detail"]:
            if row["violations"]:
                for x in row["violations"]:
                    print(f"    p{row['slide']}<{x['gate']}> {x['detail']}")
    if args.json:
        json.dump(results, open(args.json, "w"), indent=1)
        print("\nwrote", args.json)


if __name__ == "__main__":
    main()
